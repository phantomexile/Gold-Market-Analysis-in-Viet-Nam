#!/usr/bin/env python3
"""Tạo hình và bản Word cho Report 3 từ bản Markdown đã được kiểm tra số liệu."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
REPORTS = ROOT / "docs" / "reports"
MARKDOWN = REPORTS / "Report_3_Model_Results.md"
WORD = REPORTS / "Report_3_Model_Results.docx"
FIGURES = REPORTS / "figures"
FONT_FILE = Path("C:/Windows/Fonts/arial.ttf")
BLUE = "#24527A"
TEAL = "#20A39E"
ORANGE = "#E8873A"
GRAY = "#607080"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    """Chọn phông Unicode để chữ tiếng Việt trong hình không bị mất dấu."""
    path = Path("C:/Windows/Fonts/arialbd.ttf") if bold else FONT_FILE
    return ImageFont.truetype(str(path), size)


def chart_base(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    """Tạo nền và tiêu đề thống nhất cho ba hình nghiên cứu."""
    image = Image.new("RGB", (1400, 720), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 50), title, font=font(34, True), fill=BLUE)
    draw.text((80, 105), subtitle, font=font(21), fill=GRAY)
    draw.line((80, 650, 1320, 650), fill="#D6DEE5", width=2)
    return image, draw


def draw_snapshot() -> None:
    """So sánh lợi suất điểm và phân vị rủi ro của snapshot mới nhất."""
    image, draw = chart_base("Dự báo lợi suất sau spread", "Snapshot 06/10/2026 | Đơn vị: %")
    labels = ["21 quan sát", "63 quan sát", "105 quan sát"]
    values = [(-4.29, -5.41), (-4.06, -5.83), (-11.15, -13.27)]
    left, right = 250, 1260
    zero = left
    scale = (right - left) / 15
    for idx, (label, (point, q10)) in enumerate(zip(labels, values)):
        y = 205 + idx * 145
        draw.text((80, y + 18), label, font=font(23, True), fill=BLUE)
        for offset, value, color, name in [(0, point, TEAL, "Dự báo"), (48, q10, ORANGE, "q10")]:
            end = zero + abs(value) * scale
            draw.rounded_rectangle((zero, y + offset, end, y + offset + 31), radius=7, fill=color)
            draw.text((end + 15, y + offset + 1), f"{name}: {value:.2f}%", font=font(22), fill=BLUE)
    draw.text((80, 667), "Giá trị âm biểu thị lợi suất dự báo thấp hơn giá vốn mua vào.", font=font(18), fill=GRAY)
    image.save(FIGURES / "report3_snapshot_predictions.png")


def draw_mae() -> None:
    """Vẽ sai số backtest trên các mô hình có cùng số ngày đánh giá."""
    image, draw = chart_base("MAE trong tập kiểm tra tham chiếu", "Lần chạy 15/07/2026 | Đơn vị: điểm phần trăm lợi suất")
    labels = ["21", "63", "105"]
    baseline = [3.99, 11.00, 17.03]
    lightgbm = [10.15, 15.71, 11.00]
    bottom, top = 590, 190
    draw.line((120, bottom, 1270, bottom), fill=GRAY, width=2)
    for idx, label in enumerate(labels):
        x = 280 + idx * 390
        for offset, value, color in [(-76, baseline[idx], BLUE), (25, lightgbm[idx], TEAL)]:
            height = value / 20 * (bottom - top)
            draw.rectangle((x + offset, bottom - height, x + offset + 70, bottom), fill=color)
            draw.text((x + offset - 4, bottom - height - 32), f"{value:.2f}", font=font(21), fill=BLUE)
        draw.text((x - 20, 603), f"{label} quan sát", font=font(21), fill=BLUE)
    draw.rectangle((80, 145, 110, 175), fill=BLUE)
    draw.text((125, 147), "Trung bình lịch sử", font=font(20), fill=BLUE)
    draw.rectangle((420, 145, 450, 175), fill=TEAL)
    draw.text((465, 147), "LightGBM", font=font(20), fill=BLUE)
    image.save(FIGURES / "report3_test_mae.png")


def draw_signal_rate() -> None:
    """Thể hiện độ chọn lọc của quy tắc mua giữa hai pha backtest."""
    image, draw = chart_base("Tỷ lệ ngày phát tín hiệu mua", "Quy tắc mặc định trong backtest 15/07/2026 | Đơn vị: %")
    labels = ["21", "63", "105"]
    validation = [0.0, 38.7, 50.1]
    test = [4.9, 30.2, 99.3]
    bottom, top = 590, 190
    draw.line((120, bottom, 1270, bottom), fill=GRAY, width=2)
    for idx, label in enumerate(labels):
        x = 280 + idx * 390
        for offset, value, color in [(-76, validation[idx], BLUE), (25, test[idx], ORANGE)]:
            height = value / 100 * (bottom - top)
            draw.rectangle((x + offset, bottom - height, x + offset + 70, bottom), fill=color)
            draw.text((x + offset - 4, bottom - height - 31), f"{value:.1f}", font=font(21), fill=BLUE)
        draw.text((x - 20, 603), f"{label} quan sát", font=font(21), fill=BLUE)
    draw.rectangle((80, 145, 110, 175), fill=BLUE)
    draw.text((125, 147), "Xác thực", font=font(20), fill=BLUE)
    draw.rectangle((300, 145, 330, 175), fill=ORANGE)
    draw.text((345, 147), "Kiểm tra", font=font(20), fill=BLUE)
    image.save(FIGURES / "report3_signal_rate.png")


def plain_text(value: str) -> str:
    """Bỏ ký hiệu Markdown khi chuyển sang Word, giữ nguyên nội dung có dấu."""
    return value.replace("**", "").replace("`", "")


def shade(cell, fill: str) -> None:
    """Tô nền tiêu đề bảng để người đọc nhận ra tên cột trên bản in."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, lines: list[str]) -> None:
    """Chuyển một bảng Markdown thành bảng Word có hàng tiêu đề rõ ràng."""
    rows = [[plain_text(item.strip()) for item in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", item) for item in row)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            if i == 0:
                shade(cell, "24527A")
            for paragraph in cell.paragraphs:
                paragraph.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(8.3)
                    if i == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()


def add_paragraph(doc: Document, line: str) -> None:
    """Chuyển đoạn văn, đề mục, danh sách và hình sang tài liệu Word."""
    if line.startswith("!["):
        match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
        if match:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(REPORTS / match.group(1)), width=Cm(15.8))
        return
    if line.startswith("# "):
        doc.add_heading(plain_text(line[2:]), level=0)
    elif line.startswith("## "):
        doc.add_heading(plain_text(line[3:]), level=1)
    elif line.startswith("### "):
        doc.add_heading(plain_text(line[4:]), level=2)
    elif re.match(r"^\d+\. ", line):
        doc.add_paragraph(plain_text(re.sub(r"^\d+\. ", "", line)), style="List Number")
    elif line.startswith("- "):
        doc.add_paragraph(plain_text(line[2:]), style="List Bullet")
    else:
        doc.add_paragraph(plain_text(line.strip()))


def build_word() -> None:
    """Tạo bản Word từ Markdown để hai định dạng luôn có cùng kết luận."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 19), ("Heading 1", 14), ("Heading 2", 11.5)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(36, 82, 122)
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(doc, table_lines)
            continue
        add_paragraph(doc, line)
        index += 1
    doc.save(WORD)


def main() -> None:
    """Sinh đúng ba hình và một tài liệu Word từ nội dung Report 3."""
    FIGURES.mkdir(parents=True, exist_ok=True)
    draw_snapshot()
    draw_mae()
    draw_signal_rate()
    build_word()
    print(WORD)


if __name__ == "__main__":
    main()
